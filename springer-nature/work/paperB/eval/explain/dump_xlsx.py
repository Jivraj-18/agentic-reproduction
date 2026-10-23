import openpyxl, sys
for f in ["source/data/Data Fig 1+Tab 2.xlsx","source/data/Data Fig 2+Tab 2.xlsx"]:
    wb=openpyxl.load_workbook(f,data_only=True)
    for ws in wb.worksheets:
        print("=====",f,ws.title,ws.max_row,ws.max_column)
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                print("\t".join("" if v is None else (f"{v:.4g}" if isinstance(v,float) else str(v)) for v in row))
