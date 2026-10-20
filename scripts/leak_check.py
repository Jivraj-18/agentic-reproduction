"""Audit 3 (leak check): grep every file inside each sandbox for every published value.

Tiers
  inputs : files the pipeline wrote for the reproducer (TASK.md, methods.json, template.json).
           Must have zero hits.
  data   : the authors' deposited data. Short numbers ("77", "0.5") also occur as raw
           measurements, so hits here are listed for inspection, not counted as leaks.
  answers: (Paper B) full-precision fitted values from the authors' Prism tabs that were
           withheld. Must have zero hits in the sandbox: proves the answer tabs were stripped.
The reproducer's own out/ folder is excluded (it is supposed to contain its estimates);
hard-coded literals there are caught by the hardcoding audit.

Usage: python -I leak_check.py <paper_dir> <out.json> [answers_xlsx]
"""
import json
import re
import sys
from pathlib import Path


def tokens_for(cell):
    raw = cell["raw"].lstrip("<>").replace("−", "-").lstrip("-")
    forms = {raw}
    if "." in raw:  # '0.90' is also findable as '0.9'
        forms.add(raw.rstrip("0").rstrip("."))
    # A single digit ('1', '5') occurs in any text ('CB1', '0.1 nM'); it cannot be searched meaningfully.
    return [f for f in forms if sum(ch.isdigit() for ch in f) >= 2]


def unsearchable(full):
    return sorted(k for k, c in full["cells"].items() if c.get("raw") and c.get("value") is not None and not tokens_for(c))


def answer_values(xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    vals = set()
    for ws in wb.worksheets:
        if "Table" not in ws.title:
            continue
        for row in ws.iter_rows(max_row=32, values_only=True):
            for v in row:
                if isinstance(v, float) and v != 0 and len(f"{abs(v):.10g}".replace(".", "").lstrip("0")) >= 4:
                    vals.add(f"{abs(v):.10g}")
    return sorted(vals)


def scan(text, needle):
    return [m.start() for m in re.finditer(rf"(?<![\d.]){re.escape(needle)}(?!\d)", text)]


def main(paper_dir, out_path, answers_xlsx=None):
    paper_dir = Path(paper_dir)
    full = json.load(open(paper_dir / "extract" / "tables_full.json", encoding="utf-8"))
    needles = {}
    for key, c in full["cells"].items():
        if c.get("raw") and c.get("value") is not None:
            for t in tokens_for(c):
                needles.setdefault(t, []).append(key)
    answers = answer_values(answers_xlsx) if answers_xlsx else []
    report = {"published_tokens": len(needles), "answer_values": len(answers),
              "unsearchable_cells": unsearchable(full), "sandboxes": {}}
    for sb in sorted((paper_dir / "sandbox").glob("r*")):
        res = {"inputs": [], "data": [], "answers": []}
        for f in sb.rglob("*"):
            if not f.is_file() or "out" in f.relative_to(sb).parts:
                continue
            if f.suffix.lower() in (".xlsx",):
                import openpyxl
                wb = openpyxl.load_workbook(f, data_only=True)
                text = "\n".join(",".join("" if v is None else str(v) for v in r)
                                 for ws in wb.worksheets for r in ws.iter_rows(values_only=True))
            else:
                text = f.read_text(encoding="utf-8", errors="replace")
            tier = "data" if "data" in f.relative_to(sb).parts else "inputs"
            for t, keys in needles.items():
                n = len(scan(text, t))
                if n:
                    res[tier].append({"file": str(f.relative_to(sb)), "token": t, "cells": keys, "hits": n})
            for a in answers:
                n = len(scan(text, a))
                if n:
                    res["answers"].append({"file": str(f.relative_to(sb)), "value": a, "hits": n})
        res["inputs_hits"] = sum(x["hits"] for x in res["inputs"])
        res["answers_hits"] = sum(x["hits"] for x in res["answers"])
        res["data_hits"] = sum(x["hits"] for x in res["data"])
        report["sandboxes"][sb.name] = res
        print(f"{sb.name}: inputs={res['inputs_hits']} answers={res['answers_hits']} data(coincidental)={res['data_hits']}")
        for x in res["inputs"]:
            print("   INPUT HIT", x)
    report["pass"] = all(v["inputs_hits"] == 0 and v["answers_hits"] == 0 for v in report["sandboxes"].values())
    json.dump(report, open(out_path, "w"), indent=1)
    print("PASS" if report["pass"] else "FAIL")


if __name__ == "__main__":
    main(*sys.argv[1:])
